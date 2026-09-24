#!/usr/bin/env python3
"""Genera entregables/VeriTech-resultados.xlsx: el libro de resultados de la prueba del colegio.

El libro no trae datos ni números precalculados: el equipo pega los CSV de Netlify Forms en las
hojas «Pegar …» y todo lo demás se recalcula con fórmulas de Excel (sin macros ni Power Query).
Las reglas replican veritech-sitio/js/analisis.js, que es la fuente de verdad.

Uso (openpyxl vive en un entorno aislado):
    python3 -m venv ~/.venvs/veritech && ~/.venvs/veritech/bin/pip install openpyxl
    ~/.venvs/veritech/bin/python herramientas/crear_excel.py [--salida RUTA.xlsx]
"""
import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as letra
from openpyxl.worksheet.properties import PageSetupProperties

RAIZ = Path(__file__).resolve().parent.parent
SALIDA = RAIZ / "entregables" / "VeriTech-resultados.xlsx"

FILAS = 1000                 # filas de respuestas que cubren las hojas limpias
F0, F1 = 3, 3 + FILAS - 1    # fila 1: títulos · fila 2: columna encontrada · 3..1002: datos
PEGAR_DATOS = "$A$1:$BZ$1001"
PEGAR_TITULOS = "$A$1:$BZ$1"

# ---------------------------------------------------------------- estilo
AZUL = "1F3A5F"
F_TITULO = Font(bold=True, size=16, color=AZUL)
F_SECCION = Font(bold=True, size=13, color="FFFFFF")
F_NEGRITA = Font(bold=True)
F_GRIS = Font(italic=True, size=9, color="7F7F7F")
R_SECCION = PatternFill("solid", fgColor=AZUL)
R_CABECERA = PatternFill("solid", fgColor="DCE6F1")
R_AUX = PatternFill("solid", fgColor="F2F2F2")
R_VERDE = PatternFill("solid", fgColor="C6EFCE")
R_ROJO = PatternFill("solid", fgColor="FFC7CE")
R_GRIS = PatternFill("solid", fgColor="E7E6E6")
R_DIAGONAL = PatternFill("solid", fgColor="E2EFDA")
LINEA = Side(style="thin", color="BFBFBF")
BORDE = Border(top=LINEA, bottom=LINEA, left=LINEA, right=LINEA)
AJUSTE = Alignment(wrap_text=True, vertical="top")
PCT = "0%"


# ---------------------------------------------------------------- piezas de fórmula
ACENTOS = [("á", "a"), ("é", "e"), ("í", "i"), ("ó", "o"), ("ú", "u"), ("ü", "u"), ("ñ", "n"),
           ("à", "a"), ("è", "e"), ("ì", "i"), ("ò", "o"), ("ù", "u")]


def norm(x):
    """Como norm() de analisis.js: minúsculas, sin tildes y sin espacios sobrantes."""
    f = f"LOWER(TRIM({x}))"
    for con, sin in ACENTOS:
        f = f'SUBSTITUTE({f},"{con}","{sin}")'
    return f


def norm_py(s):
    s = " ".join(s.lower().split())
    for con, sin in ACENTOS:
        s = s.replace(con, sin)
    return s


def literal(x):
    """Criterio de COUNTIF que compara igualdad exacta con un texto de las respuestas:
    anula los comodines * ? ~ y los operadores (<, >, =) que pudiera traer el texto."""
    return f'"="&SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({x},"~","~~"),"*","~*"),"?","~?")'


def es_digito(x, pos):
    return f'ISNUMBER(FIND("#"&MID({x},{pos},1)&"#","#0#1#2#3#4#5#6#7#8#9#"))'


def escala(x):
    """Número inicial de la escala («5 - Muy satisfecho» → 5), como parseInt; solo vale 1 a 5."""
    dos, uno = f"VALUE(LEFT({x},2))", f"VALUE(LEFT({x},1))"
    return (f'IF({es_digito(x, 1)},IF({es_digito(x, 2)},IF(AND({dos}>=1,{dos}<=5),{dos},""),'
            f'IF(AND({uno}>=1,{uno}<=5),{uno},"")),"")')


# ---------------------------------------------------------------- hojas limpias
def variantes(nombre):
    """Nombres con los que puede llegar una columna en la fila de títulos del CSV."""
    v = [nombre]
    base = nombre[:-2] if nombre.endswith("[]") else nombre
    v += [base, base + "[]"] if nombre.endswith("[]") else [base + "[]"]
    if "_" in base:
        v.append(base.replace("_", " "))
    return list(dict.fromkeys(v))


def hoja_limpia(wb, nombre, pegar, campos, auxiliares):
    """campos: nombres de columna del CSV, en orden fijo.
    auxiliares: [(clave, título, función(c, cum, r) → fórmula sin «=»)]
    c(clave) → celda de esa fila; cum(clave) → rango desde la primera fila de datos hasta esta."""
    ws = wb.create_sheet(nombre)
    col = {k: letra(i) for i, k in enumerate(campos, 1)}
    for j, (k, _, _) in enumerate(auxiliares, len(campos) + 1):
        col[k] = letra(j)
    hoja_pegar = f"'{pegar}'"

    for k in campos:
        c = col[k]
        ws[f"{c}1"] = k
        ws[f"{c}1"].fill = R_CABECERA
        busca = "0"
        for v in reversed(variantes(k)):
            busca = f'IFERROR(MATCH("{v}",{hoja_pegar}!{PEGAR_TITULOS},0),{busca})'
        ws[f"{c}2"] = "=" + busca
        ws[f"{c}2"].number_format = '"columna "0;;"no está en el CSV"'
        ws[f"{c}2"].font = F_GRIS
        ws.column_dimensions[c].width = 18
    for k, titulo, _ in auxiliares:
        c = col[k]
        ws[f"{c}1"] = titulo
        ws[f"{c}1"].fill = R_GRIS
        ws[f"{c}2"] = "calculada"
        ws[f"{c}2"].font = F_GRIS
        ws.column_dimensions[c].width = 16
    for c in col.values():
        ws[f"{c}1"].font = F_NEGRITA
        ws[f"{c}1"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 45

    for r in range(F0, F1 + 1):
        for k in campos:
            c = col[k]
            ws[f"{c}{r}"] = (f'=IF({c}$2=0,"",IFERROR(TRIM(INDEX({hoja_pegar}!{PEGAR_DATOS},'
                             f'{r - 1},{c}$2)&""),""))')
        celda = lambda k, r=r: f"{col[k]}{r}"
        cum = lambda k, r=r: f"{col[k]}${F0}:{col[k]}{r}"
        for k, _, f in auxiliares:
            ws[f"{col[k]}{r}"] = "=" + f(celda, cum, r)
            ws[f"{col[k]}{r}"].fill = R_AUX

    ws.freeze_panes = "A3"
    ws.sheet_properties.tabColor = "5B9BD5"
    rango = lambda k: f"{nombre}!${col[k]}${F0}:${col[k]}${F1}"
    return ws, col, rango


def aparicion(valor, valida, propia):
    """N.º de orden de la primera aparición de cada valor distinto (entre filas válidas)."""
    return lambda c, cum, r: (
        f'IF(AND({c(valida)}=1,{c(valor)}<>""),IF(IFERROR(COUNTIFS({cum(valor)},{literal(c(valor))},'
        f'{cum(valida)},1),0)=1,MAX({propia}$2:{propia}{r - 1})+1,""),"")')


def orden(valor, valida, ap, rango_valor, rango_valida):
    """Puntaje para ordenar como contar(): más respuestas primero; empate → aparece antes."""
    return lambda c, cum, r: (
        f'IF({c(ap)}="","",COUNTIFS({rango_valor},{literal(c(valor))},{rango_valida},1)*10000-{c(ap)})')


def abs_local(col, k):
    return f"${col[k]}${F0}:${col[k]}${F1}"


# ---------------------------------------------------------------- datos de referencia
CAMPOS_CAL = ["docente", "area", "url", "juicio_docente", "puntaje_veritech", "nivel_veritech", "coincide",
              "senales", "palabras", "version_motor", "motivo", "consentimiento"]
CAMPOS_DOC = ["area", "grado", "funciones_probadas[]", "utilidad_semaforo", "coincide_criterio",
              "ejemplo_desacuerdo", "claridad", "usaria_en_clase", "imprescindible[]", "imprescindible_otro",
              "cambio_principal", "contratar", "precio_razonable", "preocupaciones", "correo_contacto",
              "consentimiento"]
CAMPOS_EST = ["satisfaccion", "facilidad", "funcionamiento", "funcion_mas_util", "entendio_semaforo",
              "problemas", "mejora_prioritaria", "recomendacion", "usaria_en_tarea", "grado", "consentimiento"]

SENALES = ["Autoría", "Fecha", "Referencias", "Dominio", "Escritura de relleno", "Sensacionalismo",
           "Certezas sin respaldo", "Precisión", "Profundidad"]
IMPRESCINDIBLE = ["Ver las fuentes que entrega cada alumno", "Informe del curso", "Extension instalada",
                  "Mas formatos de cita", "Integracion Classroom o Teams", "Capacitacion", "Fuentes peruanas"]

# valor guardado (sin tildes) → cómo se muestra
EQUIVALENCIAS = [
    ("confiable", "Confiable"), ("dudosa", "Dudosa"), ("no confiable", "No confiable"),
    ("Si", "Sí"), ("Tal vez", "Tal vez"), ("No", "No"), ("No me corresponde", "No me corresponde"),
    ("Casi siempre", "Casi siempre"), ("A veces", "A veces"), ("Pocas veces", "Pocas veces"),
    ("No lo compare", "No lo comparé"),
    ("Si, solos", "Sí, solos"), ("Con una explicacion breve", "Con una explicación breve"),
    ("Si, en la proxima tarea", "Sí, en la próxima tarea"), ("Si, con cambios", "Sí, con cambios"),
    ("No por ahora", "No por ahora"),
    ("Muy facil", "Muy fácil"), ("Facil", "Fácil"), ("Ni facil ni dificil", "Ni fácil ni difícil"),
    ("Dificil", "Difícil"), ("Muy dificil", "Muy difícil"),
    ("Si, todo funciono", "Sí, todo funcionó"), ("Funciono con detalles", "Funcionó con detalles"),
    ("Tuve problemas", "Tuve problemas"), ("No pude usarla", "No pude usarla"),
    ("Si, lo entendi", "Sí, lo entendí"), ("Mas o menos", "Más o menos"), ("No lo entendi", "No lo entendí"),
    ("No analice ninguna fuente", "No analicé ninguna fuente"),
    ("Ver las fuentes que entrega cada alumno", "Ver las fuentes que entrega cada alumno"),
    ("Informe del curso", "Informe del curso"),
    ("Extension instalada", "Que funcione solo en el navegador (extensión)"),
    ("Mas formatos de cita", "Más formatos de cita"),
    ("Integracion Classroom o Teams", "Integración con Classroom o Teams"),
    ("Capacitacion", "Capacitación o guía para docentes"),
    ("Fuentes peruanas", "Que reconozca fuentes peruanas"),
]


# ---------------------------------------------------------------- libro
def construir():
    wb = Workbook()
    ins = wb.active
    ins.title = "Instrucciones"
    for nombre in ("Pegar calibracion", "Pegar docentes", "Pegar estudiantes"):
        p = wb.create_sheet(nombre)
        p.freeze_panes = "A2"
        p.sheet_properties.tabColor = "FFC000"

    # --- Calibracion
    cal_aux = [
        ("jn", "juicio docente (normalizado)", lambda c, cum, r: f'IF({c("juicio_docente")}="","",{norm(c("juicio_docente"))})'),
        ("vn", "nivel VeriTech (normalizado)", lambda c, cum, r: f'IF({c("nivel_veritech")}="","",{norm(c("nivel_veritech"))})'),
        ("jnum", "Nivel docente (2 confiable · 1 dudosa · 0 no confiable)",
         lambda c, cum, r: f'IF({c("jn")}="confiable",2,IF({c("jn")}="dudosa",1,IF({c("jn")}="no confiable",0,"")))'),
        ("vnum", "Nivel VeriTech (2 confiable · 1 dudosa · 0 no confiable)",
         lambda c, cum, r: f'IF({c("vn")}="confiable",2,IF({c("vn")}="dudosa",1,IF({c("vn")}="no confiable",0,"")))'),
        ("valida", "Fila válida (1 = tiene los dos niveles)",
         lambda c, cum, r: f'IF(AND(ISNUMBER({c("jnum")}),ISNUMBER({c("vnum")})),1,"")'),
        ("coinc", "¿Coincide? (Sí/No)",
         lambda c, cum, r: f'IF({c("valida")}=1,IF({c("jnum")}={c("vnum")},"Sí","No"),"")'),
        ("dir", "VeriTech fue…",
         lambda c, cum, r: f'IF({c("valida")}=1,IF({c("vnum")}<{c("jnum")},"Más estricto",'
                           f'IF({c("vnum")}>{c("jnum")},"Más permisivo","Igual")),"")'),
        ("grave", "Desacuerdo grave (1 = confiable vs no confiable)",
         lambda c, cum, r: f'IF({c("valida")}=1,IF(ABS({c("jnum")}-{c("vnum")})=2,1,0),"")'),
        ("sen", "señales (normalizadas)",
         lambda c, cum, r: f'IF({c("senales")}="","","|"&SUBSTITUTE(SUBSTITUTE(TRIM({c("senales")})," |","|"),"| ","|")&"|")'),
    ]
    # columnas que dependen de la letra propia: se completan tras conocerla
    cal_aux += [("ap_area", "Área: n.º de aparición", None), ("ord_area", "Área: orden", None),
                ("doc_clave", "docente (clave)", None), ("doc_nuevo", "Docente nuevo (1)", None),
                ("grave_n", "Grave n.º", None)]
    letras_cal = {k: letra(i) for i, k in enumerate(CAMPOS_CAL + [a[0] for a in cal_aux], 1)}
    cal_aux[-5] = ("ap_area", "Área: n.º de aparición", aparicion("area", "valida", letras_cal["ap_area"]))
    cal_aux[-4] = ("ord_area", "Área: orden", orden("area", "valida", "ap_area",
                                                     abs_local(letras_cal, "area"), abs_local(letras_cal, "valida")))
    cal_aux[-3] = ("doc_clave", "docente (clave)",
                   lambda c, cum, r: f'IF({c("valida")}=1,IF({norm(c("docente"))}="","(sin apodo)",{norm(c("docente"))}),"")')
    cal_aux[-2] = ("doc_nuevo", "Docente nuevo (1)",
                   lambda c, cum, r: f'IF({c("doc_clave")}="","",IF(IFERROR(COUNTIF({cum("doc_clave")},{literal(c("doc_clave"))}),0)=1,1,""))')
    g = letras_cal["grave_n"]
    cal_aux[-1] = ("grave_n", "Grave n.º",
                   lambda c, cum, r: f'IF({c("grave")}=1,MAX({g}$2:{g}{r - 1})+1,"")')
    _, _, CAL = hoja_limpia(wb, "Calibracion", "Pegar calibracion", CAMPOS_CAL, cal_aux)

    # --- Docentes
    doc_aux = [
        ("valida", "Fila válida (1 = respondió utilidad o uso en clase)",
         lambda c, cum, r: f'IF(OR({c("utilidad_semaforo")}<>"",{c("usaria_en_clase")}<>""),1,"")'),
        ("util", "Utilidad del semáforo (1 a 5)",
         lambda c, cum, r: f'IF({c("valida")}=1,{escala(c("utilidad_semaforo"))},"")'),
    ]
    for campo, titulo in [("coincide_criterio", "coincide_criterio"), ("claridad", "claridad"),
                          ("usaria_en_clase", "usaria_en_clase"), ("contratar", "contratar"),
                          ("imprescindible[]", "imprescindible")]:
        doc_aux.append(("n_" + campo, titulo + " (normalizado)",
                        lambda c, cum, r, campo=campo: f'IF(AND({c("valida")}=1,{c(campo)}<>""),{norm(c(campo))},"")'))
    doc_aux.append(("usa_si", "Usaría en clase (1 = «Si, en la proxima tarea» o «Si, con cambios»)",
                    lambda c, cum, r: f'IF({c("n_usaria_en_clase")}="","",IF(OR({c("n_usaria_en_clase")}="si, en la proxima tarea",'
                                      f'{c("n_usaria_en_clase")}="si, con cambios"),1,0))'))
    doc_aux += [("ap_area", "Área: n.º de aparición", None), ("ord_area", "Área: orden", None)]
    letras_doc = {k: letra(i) for i, k in enumerate(CAMPOS_DOC + [a[0] for a in doc_aux], 1)}
    doc_aux[-2] = ("ap_area", "Área: n.º de aparición", aparicion("area", "valida", letras_doc["ap_area"]))
    doc_aux[-1] = ("ord_area", "Área: orden", orden("area", "valida", "ap_area",
                                                     abs_local(letras_doc, "area"), abs_local(letras_doc, "valida")))
    _, _, DOC = hoja_limpia(wb, "Docentes", "Pegar docentes", CAMPOS_DOC, doc_aux)

    # --- Estudiantes
    est_aux = [
        ("valida", "Fila válida (1 = respondió satisfacción)",
         lambda c, cum, r: f'IF({c("satisfaccion")}<>"",1,"")'),
        ("satis", "Satisfacción (1 a 5)", lambda c, cum, r: f'IF({c("valida")}=1,{escala(c("satisfaccion"))},"")'),
        ("recom", "Recomendación (1 a 5)", lambda c, cum, r: f'IF({c("valida")}=1,{escala(c("recomendacion"))},"")'),
    ]
    for campo in ("facilidad", "funcionamiento", "entendio_semaforo", "usaria_en_tarea"):
        est_aux.append(("n_" + campo, campo + " (normalizado)",
                        lambda c, cum, r, campo=campo: f'IF(AND({c("valida")}=1,{c(campo)}<>""),{norm(c(campo))},"")'))
    est_aux.append(("usa_si", "Usaría en su tarea (1 = «Si» o «Tal vez»)",
                    lambda c, cum, r: f'IF({c("n_usaria_en_tarea")}="","",IF(OR({c("n_usaria_en_tarea")}="si",'
                                      f'{c("n_usaria_en_tarea")}="tal vez"),1,0))'))
    est_aux += [("ap_fun", "Función: n.º de aparición", None), ("ord_fun", "Función: orden", None),
                ("ap_grado", "Grado: n.º de aparición", None), ("ord_grado", "Grado: orden", None)]
    letras_est = {k: letra(i) for i, k in enumerate(CAMPOS_EST + [a[0] for a in est_aux], 1)}
    est_aux[-4] = ("ap_fun", "Función: n.º de aparición", aparicion("funcion_mas_util", "valida", letras_est["ap_fun"]))
    est_aux[-3] = ("ord_fun", "Función: orden", orden("funcion_mas_util", "valida", "ap_fun",
                                                       abs_local(letras_est, "funcion_mas_util"), abs_local(letras_est, "valida")))
    est_aux[-2] = ("ap_grado", "Grado: n.º de aparición", aparicion("grado", "valida", letras_est["ap_grado"]))
    est_aux[-1] = ("ord_grado", "Grado: orden", orden("grado", "valida", "ap_grado",
                                                       abs_local(letras_est, "grado"), abs_local(letras_est, "valida")))
    _, _, EST = hoja_limpia(wb, "Estudiantes", "Pegar estudiantes", CAMPOS_EST, est_aux)

    celdas = resumen(wb, CAL, DOC, EST)
    instrucciones(ins)

    for ws in wb.worksheets:
        if not ws.title.startswith("Pegar"):
            ws.protection.sheet = True
            ws.protection.formatColumns = False
            ws.protection.formatRows = False
            ws.protection.formatCells = False
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    wb.move_sheet("Resumen", offset=-(len(wb.sheetnames) - 2))
    wb.calculation.fullCalcOnLoad = True
    return wb, celdas


# ---------------------------------------------------------------- Resumen
def resumen(wb, CAL, DOC, EST):
    ws = wb.create_sheet("Resumen")
    ws.sheet_properties.tabColor = "70AD47"
    for c, w in zip("ABCDEF", (58, 14, 14, 14, 40, 26)):
        ws.column_dimensions[c].width = w
    ws.column_dimensions["F"].hidden = True   # claves internas de cada fila
    celdas = {}
    fila = [1]
    graficos = []

    def poner(ref, valor, fmt=None, font=None, fill=None):
        ws[ref] = valor
        if fmt:
            ws[ref].number_format = fmt
        if font:
            ws[ref].font = font
        if fill:
            ws[ref].fill = fill
        return ref

    def seccion(texto):
        r = fila[0] + 1
        for c in "ABCDE":
            ws[f"{c}{r}"].fill = R_SECCION
        poner(f"A{r}", texto, font=F_SECCION)
        fila[0] = r + 1

    def cabecera(*titulos):
        r = fila[0]
        for c, t in zip("ABCDE", titulos):
            poner(f"{c}{r}", t, font=F_NEGRITA, fill=R_CABECERA)
            ws[f"{c}{r}"].alignment = Alignment(wrap_text=True, vertical="center")
        fila[0] += 1

    def dato(clave, texto, formula, fmt=None):
        r = fila[0]
        poner(f"A{r}", texto)
        celdas[clave] = poner(f"B{r}", formula, fmt, F_NEGRITA)
        fila[0] += 1
        return f"B{r}"

    def muestra(clave):
        return f'IFERROR(VLOOKUP({clave},$A${EQ0}:$B${EQ1},2,FALSE),{clave})'

    def grafico(titulo, r_ini, r_fin, alto_filas):
        ch = BarChart()
        ch.type = "bar"
        ch.title = titulo
        ch.legend = None
        ch.style = 10
        ch.add_data(Reference(ws, min_col=2, min_row=r_ini - 1, max_row=r_fin), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=r_ini, max_row=r_fin))
        ch.x_axis.scaling.orientation = "maxMin"
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        ch.y_axis.majorGridlines = None
        ch.y_axis.number_format = "0"
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showVal = True
        ch.width, ch.height = 15, max(5.5, alto_filas * 0.5)
        graficos.append((ch, r_ini - 2))

    def distribucion(clave, titulo, rango, opciones, con_grafico=True):
        """opciones: valores guardados. Cuenta sobre la columna normalizada."""
        cabecera(titulo, "Personas", "%")
        r0 = fila[0]
        den = f"COUNTIF({rango},\"?*\")"
        for v in opciones:
            r = fila[0]
            poner(f"F{r}", v)
            poner(f"A{r}", "=" + muestra(f"F{r}"))
            celdas[f"{clave}|{v}|n"] = poner(f"B{r}", f'=COUNTIF({rango},"={norm_py(v)}")')
            celdas[f"{clave}|{v}|pct"] = poner(f"C{r}", f'=IF({den}=0,"",ROUND(B{r}/{den}*100,0)/100)', PCT)
            fila[0] += 1
        r = fila[0]
        poner(f"A{r}", "Otra respuesta", font=F_GRIS)
        poner(f"B{r}", f"={den}-SUM(B{r0}:B{r - 1})")
        poner(f"C{r}", f'=IF({den}=0,"",ROUND(B{r}/{den}*100,0)/100)', PCT)
        fila[0] += 1
        r = fila[0]
        poner(f"A{r}", "Total con respuesta", font=F_NEGRITA)
        celdas[f"{clave}|total"] = poner(f"B{r}", f"={den}", font=F_NEGRITA)
        fila[0] += 1
        if con_grafico:
            grafico(titulo, r0, r0 + len(opciones) - 1, len(opciones) + 3)
            fila[0] = max(fila[0], r0 + 12)
        fila[0] += 1

    def escala_dist(clave, titulo, rango):
        cabecera(titulo, "Personas", "%")
        r0 = fila[0]
        for k in (5, 4, 3, 2, 1):
            r = fila[0]
            poner(f"A{r}", str(k))
            ws[f"A{r}"].alignment = Alignment(horizontal="left")
            celdas[f"{clave}|{k}|n"] = poner(f"B{r}", f"=COUNTIF({rango},{k})")
            poner(f"C{r}", f'=IF(COUNT({rango})=0,"",ROUND(B{r}/COUNT({rango})*100,0)/100)', PCT)
            fila[0] += 1
        grafico(titulo, r0, r0 + 4, 8)
        fila[0] = max(fila[0], r0 + 12) + 1

    def dinamica(clave, titulo, rango_valor, rango_orden, rango_valida, filas=10):
        """Valores que no se conocen de antemano (área, grado…), ordenados como contar()."""
        cabecera(titulo, "Personas", "%")
        den = f'COUNTIFS({rango_valida},1,{rango_valor},"?*")'
        for k in range(1, filas + 1):
            r = fila[0]
            poner(f"A{r}", f'=IFERROR(INDEX({rango_valor},MATCH(LARGE({rango_orden},{k}),{rango_orden},0)),"")')
            celdas[f"{clave}|#{k}|valor"] = f"A{r}"
            celdas[f"{clave}|#{k}|n"] = poner(f"B{r}", f'=IF(A{r}="","",ROUNDUP(LARGE({rango_orden},{k})/10000,0))')
            celdas[f"{clave}|#{k}|pct"] = poner(f"C{r}", f'=IF(A{r}="","",ROUND(B{r}/{den}*100,0)/100)', PCT)
            fila[0] += 1
        fila[0] += 1

    # la tabla de equivalencias va al final, pero sus filas se fijan ya
    EQ0 = 400
    EQ1 = EQ0 + len(EQUIVALENCIAS) - 1

    # ------------------------------------------------ cabecera y criterios
    poner("A1", "VeriTech · Resultados de la prueba del colegio", font=F_TITULO)
    poner("A2", "Esta hoja se calcula sola con lo que pegues en las hojas «Pegar …». No escribas aquí.", font=F_GRIS)
    fila[0] = 3
    seccion("¿Pedimos el piloto? Los 4 criterios")
    cabecera("Criterio", "Valor", "Meta", "¿Se cumple?", "Detalle")

    cal_n = f"COUNT({CAL('valida')})"
    doc_n = f"COUNT({DOC('valida')})"
    est_n = f"COUNT({EST('valida')})"
    r = fila[0]
    rc1, rc2, rc3, rc4 = r, r + 1, r + 2, r + 3
    poner(f"A{rc1}", "1. El semáforo coincide con los docentes en al menos el 70 % de las fuentes")
    celdas["crit1"] = poner(f"B{rc1}", f'=IF({cal_n}=0,"",ROUND(COUNTIF({CAL("coinc")},"Sí")/{cal_n}*100,0)/100)', PCT)
    poner(f"C{rc1}", "≥ 70 %")
    poner(f"D{rc1}", f'=IF(B{rc1}="","sin datos",IF(ROUND(B{rc1}*100,0)>=70,"✓","✕"))')
    poner(f"E{rc1}", f'=IF(B{rc1}="","Falta pegar el CSV de calibración",ROUND(B{rc1}*100,0)&" % en "&{cal_n}&" fuentes")')

    poner(f"A{rc2}", "2. Utilidad del semáforo según los docentes: promedio de 4 o más (de 5)")
    celdas["crit2"] = poner(f"B{rc2}", f'=IF(COUNT({DOC("util")})=0,"",ROUND(AVERAGE({DOC("util")}),1))', "0.0")
    poner(f"C{rc2}", "≥ 4 de 5")
    poner(f"D{rc2}", f'=IF(B{rc2}="","sin datos",IF(B{rc2}>=4,"✓","✕"))')
    poner(f"E{rc2}", f'=IF({doc_n}=0,"Falta pegar el CSV de opinión docente",IF(B{rc2}="","sin respuestas",B{rc2}&" de 5"))')

    poner(f"A{rc3}", "3. Más de la mitad de los docentes lo usaría en clase («Sí, en la próxima tarea» o «Sí, con cambios»)")
    celdas["crit3"] = poner(f"B{rc3}", f'=IF(COUNT({DOC("usa_si")})=0,"",SUM({DOC("usa_si")})/COUNT({DOC("usa_si")}))', PCT)
    poner(f"C{rc3}", "más de 50 %")
    poner(f"D{rc3}", f'=IF(B{rc3}="","sin datos",IF(B{rc3}>0.5,"✓","✕"))')
    poner(f"E{rc3}", f'=IF({doc_n}=0,"Falta pegar el CSV de opinión docente",IF(B{rc3}="","sin respuestas",'
                     f'ROUND(B{rc3}*100,0)&" % de "&{doc_n}&" docentes"))')

    poner(f"A{rc4}", "4. Más de la mitad de los estudiantes lo usaría en su próxima tarea («Sí» o «Tal vez»)")
    celdas["crit4"] = poner(f"B{rc4}", f'=IF(COUNT({EST("usa_si")})=0,"",SUM({EST("usa_si")})/COUNT({EST("usa_si")}))', PCT)
    poner(f"C{rc4}", "más de 50 %")
    poner(f"D{rc4}", f'=IF(B{rc4}="","sin datos",IF(B{rc4}>0.5,"✓","✕"))')
    poner(f"E{rc4}", f'=IF({est_n}=0,"Falta pegar el CSV de estudiantes",IF(B{rc4}="","sin respuestas",'
                     f'ROUND(B{rc4}*100,0)&" % de "&{est_n}&" estudiantes"))')
    for rr in (rc1, rc2, rc3, rc4):
        ws[f"A{rr}"].alignment = AJUSTE
        ws[f"E{rr}"].alignment = AJUSTE
        ws[f"D{rr}"].alignment = Alignment(horizontal="center", vertical="top")
        ws[f"D{rr}"].font = Font(bold=True, size=13)
        ws.row_dimensions[rr].height = 32
        for c in "ABCDE":
            ws[f"{c}{rr}"].border = BORDE
        celdas[f"estado{rr - rc1 + 1}"] = f"D{rr}"
    rango_estado = f"D{rc1}:D{rc4}"
    ws.conditional_formatting.add(rango_estado, CellIsRule(operator="equal", formula=['"✓"'], fill=R_VERDE, font=Font(bold=True, color="006100")))
    ws.conditional_formatting.add(rango_estado, CellIsRule(operator="equal", formula=['"✕"'], fill=R_ROJO, font=Font(bold=True, color="9C0006")))
    ws.conditional_formatting.add(rango_estado, CellIsRule(operator="equal", formula=['"sin datos"'], fill=R_GRIS, font=Font(color="595959")))
    fila[0] = rc4 + 1
    r = fila[0]
    poner(f"A{r}", "Criterios cumplidos", font=F_NEGRITA)
    celdas["cumplidos"] = poner(f"B{r}", f'=COUNTIF({rango_estado},"✓")&" de 4"', font=F_NEGRITA)
    r += 1
    celdas["alerta"] = poner(
        f"A{r}", f'=IF(AND({cal_n}>0,B{rc1}<>""),IF(ROUND(B{rc1}*100,0)<50,"⚠ La coincidencia es menor al 50 %: '
                 f'según el plan, no se propone el piloto hasta corregir el motor y repetir la calibración.",""),"")')
    ws.merge_cells(f"A{r}:E{r}")
    ws[f"A{r}"].alignment = AJUSTE
    ws[f"A{r}"].font = Font(bold=True, color="9C0006", size=12)
    ws.row_dimensions[r].height = 34
    ws.conditional_formatting.add(f"A{r}:E{r}", FormulaRule(formula=[f'LEN($A${r})>0'], fill=R_ROJO))
    fila[0] = r + 1

    # ------------------------------------------------ calibración
    seccion("Calibración: docentes vs VeriTech")
    dato("cal_n", "Fuentes calificadas", f"={cal_n}")
    dato("cal_acuerdoN", "Fuentes en las que coinciden", f'=COUNTIF({CAL("coinc")},"Sí")')
    dato("cal_acuerdo", "Coincidencia", f"=B{rc1}", PCT)
    dato("cal_docentes", "Docentes que calificaron", f'=SUM({CAL("doc_nuevo")})')
    dato("cal_graves", "Desacuerdos graves (uno dice «confiable» y el otro «no confiable»)", f'=SUM({CAL("grave")})')
    dato("cal_estricto", "Desacuerdos en que VeriTech fue más estricto", f'=COUNTIF({CAL("dir")},"Más estricto")')
    r_estr = fila[0] - 1
    dato("cal_permisivo", "Desacuerdos en que VeriTech fue más permisivo", f'=COUNTIF({CAL("dir")},"Más permisivo")')
    r_perm = fila[0] - 1
    fila[0] += 1

    r = fila[0]
    poner(f"A{r}", "Matriz: juicio del docente (filas) vs nivel de VeriTech (columnas)", font=F_NEGRITA)
    r += 1
    niveles = ["confiable", "dudosa", "no confiable"]
    poner(f"A{r}", "Docente ↓   /   VeriTech →", font=F_NEGRITA, fill=R_CABECERA)
    for j, v in enumerate(niveles):
        c = "BCD"[j]
        poner(f"{c}{r}", v.capitalize(), font=F_NEGRITA, fill=R_CABECERA)
    for i, d in enumerate(niveles):
        rr = r + 1 + i
        poner(f"A{rr}", d.capitalize(), font=F_NEGRITA, fill=R_CABECERA)
        for j, v in enumerate(niveles):
            c = "BCD"[j]
            celdas[f"matriz|{i}|{j}"] = poner(
                f"{c}{rr}", f'=COUNTIFS({CAL("jnum")},{2 - i},{CAL("vnum")},{2 - j})',
                fill=R_DIAGONAL if i == j else None)
            ws[f"{c}{rr}"].alignment = Alignment(horizontal="center")
            ws[f"{c}{rr}"].border = BORDE
    poner(f"A{r + 4}", "En verde: coinciden. Las esquinas «Confiable / No confiable» son los desacuerdos graves.", font=F_GRIS)
    fila[0] = r + 6

    cabecera("Señal (en los desacuerdos)", "Más estricto: salió «bad»", "%", "Más permisivo: salió «ok»", "%")
    r0 = fila[0]
    for s in SENALES:
        r = fila[0]
        poner(f"A{r}", s)
        celdas[f"estricto|{s}|n"] = poner(f"B{r}", f'=COUNTIFS({CAL("dir")},"Más estricto",{CAL("sen")},"*|{s}:bad:*")')
        celdas[f"estricto|{s}|pct"] = poner(f"C{r}", f'=IF($B${r_estr}=0,"",ROUND(B{r}/$B${r_estr}*100,0)/100)', PCT)
        celdas[f"permisivo|{s}|n"] = poner(f"D{r}", f'=COUNTIFS({CAL("dir")},"Más permisivo",{CAL("sen")},"*|{s}:ok:*")')
        celdas[f"permisivo|{s}|pct"] = poner(f"E{r}", f'=IF($B${r_perm}=0,"",ROUND(D{r}/$B${r_perm}*100,0)/100)', PCT)
        ws[f"E{r}"].alignment = Alignment(horizontal="right")
        fila[0] += 1
    poner(f"A{fila[0]}", "Cuántas veces cada señal salió mal cuando VeriTech fue más estricto que el docente, "
                         "y bien cuando fue más permisivo. El % es sobre esos desacuerdos.", font=F_GRIS)
    grafico("Señales que salieron «bad» cuando VeriTech fue más estricto", r0, r0 + len(SENALES) - 1, 12)
    fila[0] = max(fila[0], r0 + 12) + 2

    r = fila[0]
    poner(f"A{r}", "Desacuerdos graves (los primeros 10)", font=F_NEGRITA)
    fila[0] += 1
    cabecera("URL", "Docente dijo", "VeriTech dijo", "Puntaje", "Motivo del docente")
    for k in range(1, 11):
        r = fila[0]
        busca = f'MATCH({k},{CAL("grave_n")},0)'
        for c, campo in zip("ABCDE", ("url", "juicio_docente", "nivel_veritech", "puntaje_veritech", "motivo")):
            valor = f"INDEX({CAL(campo)},{busca})"
            if campo in ("juicio_docente", "nivel_veritech"):
                valor = muestra(valor)
            poner(f"{c}{r}", f'=IFERROR({valor},"")')
        fila[0] += 1
    fila[0] += 1

    cabecera("Coincidencia por área", "Fuentes", "Coincidencia")
    for k in range(1, 11):
        r = fila[0]
        poner(f"A{r}", f'=IFERROR(INDEX({CAL("area")},MATCH(LARGE({CAL("ord_area")},{k}),{CAL("ord_area")},0)),"")')
        celdas[f"area|#{k}|valor"] = f"A{r}"
        celdas[f"area|#{k}|n"] = poner(f"B{r}", f'=IF(A{r}="","",ROUNDUP(LARGE({CAL("ord_area")},{k})/10000,0))')
        celdas[f"area|#{k}|acuerdo"] = poner(
            f"C{r}", f'=IF(A{r}="","",ROUND(SUMPRODUCT(({CAL("area")}=A{r})*({CAL("coinc")}="Sí"))/B{r}*100,0)/100)', PCT)
        fila[0] += 1
    fila[0] += 1

    # ------------------------------------------------ docentes
    seccion("Opinión de los docentes")
    dato("doc_n", "Docentes que respondieron", f"={doc_n}")
    dato("doc_util", "Utilidad del semáforo (promedio de 1 a 5)", f"=B{rc2}", "0.0")
    dato("doc_usaria", "Lo usarían en clase («Sí, en la próxima tarea» o «Sí, con cambios»)", f"=B{rc3}", PCT)
    fila[0] += 1
    escala_dist("doc_util", "Utilidad del semáforo (5 = muy útil)", DOC("util"))
    distribucion("coincide", "¿El semáforo coincide con su criterio?", DOC("n_coincide_criterio"),
                 ["Casi siempre", "A veces", "Pocas veces", "No lo compare"])
    distribucion("claridad", "¿Los alumnos entienden el semáforo?", DOC("n_claridad"),
                 ["Si, solos", "Con una explicacion breve", "No"])
    distribucion("usaria", "¿Lo usaría en clase?", DOC("n_usaria_en_clase"),
                 ["Si, en la proxima tarea", "Si, con cambios", "No por ahora"])
    distribucion("contratar", "¿Recomendaría que el colegio lo contrate?", DOC("n_contratar"),
                 ["Si", "Tal vez", "No", "No me corresponde"])

    cabecera("Imprescindible para usarlo (puede marcar varias)", "Docentes", "% de docentes")
    r0 = fila[0]
    for v in IMPRESCINDIBLE:
        r = fila[0]
        poner(f"F{r}", v)
        poner(f"A{r}", "=" + muestra(f"F{r}"))
        celdas[f"impr|{v}|n"] = poner(f"B{r}", f'=COUNTIF({DOC("n_imprescindible[]")},"*{norm_py(v)}*")')
        celdas[f"impr|{v}|pct"] = poner(f"C{r}", f'=IF({doc_n}=0,"",ROUND(B{r}/{doc_n}*100,0)/100)', PCT)
        fila[0] += 1
    grafico("Imprescindible para usarlo", r0, fila[0] - 1, 10)
    fila[0] = max(fila[0], r0 + 12) + 1

    dinamica("doc_area", "Área de los docentes", DOC("area"), DOC("ord_area"), DOC("valida"), 8)

    # ------------------------------------------------ estudiantes
    seccion("Opinión de los estudiantes")
    dato("est_n", "Estudiantes que respondieron", f"={est_n}")
    dato("est_satis", "Satisfacción (promedio de 1 a 5)",
         f'=IF(COUNT({EST("satis")})=0,"",ROUND(AVERAGE({EST("satis")}),1))', "0.0")
    dato("est_recom", "Recomendación (promedio de 1 a 5)",
         f'=IF(COUNT({EST("recom")})=0,"",ROUND(AVERAGE({EST("recom")}),1))', "0.0")
    dato("est_usaria", "Lo usarían en su próxima tarea («Sí» o «Tal vez»)", f"=B{rc4}", PCT)
    fila[0] += 1
    escala_dist("est_satis", "Satisfacción (5 = muy satisfecho)", EST("satis"))
    escala_dist("est_recom", "¿Lo recomendaría? (5 = muy probable)", EST("recom"))
    distribucion("facilidad", "¿Qué tan fácil fue usarla?", EST("n_facilidad"),
                 ["Muy facil", "Facil", "Ni facil ni dificil", "Dificil", "Muy dificil"])
    distribucion("funcionamiento", "¿Funcionó bien?", EST("n_funcionamiento"),
                 ["Si, todo funciono", "Funciono con detalles", "Tuve problemas", "No pude usarla"])
    distribucion("entendio", "¿Entendió el semáforo?", EST("n_entendio_semaforo"),
                 ["Si, lo entendi", "Mas o menos", "No lo entendi", "No analice ninguna fuente"])
    distribucion("est_usaria", "¿La usaría en su próxima tarea?", EST("n_usaria_en_tarea"), ["Si", "Tal vez", "No"])
    dinamica("funcion", "Función más útil", EST("funcion_mas_util"), EST("ord_fun"), EST("valida"), 8)
    dinamica("grado", "Grado", EST("grado"), EST("ord_grado"), EST("valida"), 6)

    # ------------------------------------------------ equivalencias
    assert fila[0] < EQ0 - 3, "el Resumen invade la tabla de equivalencias"
    fila[0] = EQ0 - 3
    seccion("Tabla de equivalencias: cómo se guarda cada respuesta y cómo se muestra")
    cabecera("Valor guardado (sin tildes)", "Se muestra como")
    ws.merge_cells(f"B{EQ0 - 1}:E{EQ0 - 1}")
    for i, (guardado, visible) in enumerate(EQUIVALENCIAS):
        r = EQ0 + i
        poner(f"A{r}", guardado)
        poner(f"B{r}", visible)
        ws.merge_cells(f"B{r}:E{r}")

    for ch, r in graficos:
        ws.add_chart(ch, f"H{r}")
    ws.freeze_panes = "A3"
    return celdas


# ---------------------------------------------------------------- Instrucciones
def instrucciones(ws):
    ws.sheet_properties.tabColor = AZUL
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 110
    filas = [
        ("t", "VeriTech · Cómo usar este libro de resultados"),
        ("n", "Este libro calcula solo los resultados de la prueba. Tú únicamente pegas los datos; "
              "no hace falta escribir fórmulas ni tocar las demás hojas."),
        ("s", "Paso 1 · Descargar los 3 archivos de Netlify"),
        ("1", "Entra a Netlify, abre el sitio de VeriTech y ve a Forms (Formularios)."),
        ("2", "Haz clic en el formulario «calibracion» y luego en el botón Download CSV. Se descarga un archivo .csv."),
        ("3", "Repite con «opinion-docentes» y con «opinion-sobre-la-web» (este es el de los estudiantes)."),
        ("s", "Paso 2 · Abrir cada CSV respetando las tildes (UTF-8)"),
        ("1", "En Excel, crea un libro NUEVO y en blanco (Archivo → Nuevo). No lo hagas dentro de este libro."),
        ("2", "En ese libro nuevo ve a la pestaña Datos → Desde texto/CSV y elige el archivo descargado."),
        ("3", "En «Origen del archivo» elige «65001: Unicode (UTF-8)». Revisa que palabras como «Comunicación» "
              "se vean bien, con su tilde. Luego pulsa Cargar."),
        ("·", "En Mac, si no ves «Desde texto/CSV», usa Archivo → Importar → Archivo CSV y elige "
              "«Unicode (UTF-8)» como origen del archivo."),
        ("s", "Paso 3 · Pegar los datos aquí"),
        ("1", "En el libro nuevo, haz clic en cualquier celda de la tabla y pulsa Ctrl+A (en Mac, Cmd+A) "
              "dos veces para seleccionar todo, incluida la fila de títulos. Copia con Ctrl+C (Cmd+C)."),
        ("2", "Vuelve a este libro y abre la hoja que corresponde:\n"
              "   • calibracion → hoja «Pegar calibracion»\n"
              "   • opinion-docentes → hoja «Pegar docentes»\n"
              "   • opinion-sobre-la-web → hoja «Pegar estudiantes»"),
        ("3", "Haz clic en la celda A1 (arriba a la izquierda) y pega con Ctrl+V (Cmd+V). "
              "La fila 1 tiene que ser la de los títulos (docente, area, url…)."),
        ("4", "Abre la hoja «Resumen»: todos los números y gráficos ya están actualizados."),
        ("s", "Para actualizar con datos nuevos"),
        ("1", "Descarga el CSV otra vez: Netlify siempre entrega todas las respuestas, no solo las nuevas."),
        ("2", "En la hoja «Pegar …», haz clic en el cuadrito de la esquina (arriba a la izquierda, sobre la fila 1) "
              "y pulsa Suprimir para borrar lo anterior. Después pega en A1 como en el paso 3."),
        ("s", "Qué hay en cada hoja"),
        ("·", "Resumen: los 4 criterios para pedir el piloto (✓ se cumple, ✕ no se cumple, «sin datos» falta el CSV), "
              "la calibración, la opinión de docentes y de estudiantes, con gráficos."),
        ("·", "Pegar calibracion / Pegar docentes / Pegar estudiantes: el único lugar donde se pega. "
              "Las columnas pueden venir en cualquier orden y con columnas de más: el libro las busca por su nombre."),
        ("·", "Calibracion / Docentes / Estudiantes: los datos ordenados y las columnas calculadas (en gris). "
              "La fila 2 dice en qué columna encontró cada dato o «no está en el CSV». Se llenan solas: no se escriben."),
        ("s", "Si algo no se ve bien"),
        ("·", "Letras raras en lugar de tildes (por ejemplo «ComunicaciÃ³n»): repite el paso 2 eligiendo UTF-8."),
        ("·", "Todo sale «sin datos»: revisa que pegaste en A1 de la hoja correcta y que la fila 1 tiene los títulos."),
        ("·", "Una columna dice «no está en el CSV»: descargaste otro formulario o la columna cambió de nombre en la web."),
        ("·", "Las hojas que no son «Pegar …» están protegidas para no borrar fórmulas por error "
              "(Revisar → Desproteger hoja, sin contraseña)."),
    ]
    r = 1
    for tipo, texto in filas:
        if tipo == "t":
            ws[f"A{r}"] = texto
            ws[f"A{r}"].font = F_TITULO
            r += 1
        elif tipo == "s":
            r += 1
            ws[f"A{r}"] = texto
            ws[f"A{r}"].font = F_SECCION
            ws[f"A{r}"].fill = R_SECCION
            ws[f"B{r}"].fill = R_SECCION
        elif tipo == "n":
            ws[f"A{r}"] = texto
            ws.merge_cells(f"A{r}:B{r}")
            ws[f"A{r}"].alignment = AJUSTE
            ws.row_dimensions[r].height = 32
        else:
            ws[f"A{r}"] = tipo
            ws[f"A{r}"].font = Font(bold=True, size=12, color=AZUL)
            ws[f"A{r}"].alignment = Alignment(horizontal="center", vertical="top")
            ws[f"B{r}"] = texto
            ws[f"B{r}"].alignment = AJUSTE
            lineas = texto.count("\n") + 1 + len(texto) // 105
            ws.row_dimensions[r].height = max(18, 16 * lineas)
        r += 1


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--salida", type=Path, default=SALIDA)
    args = ap.parse_args()
    wb, _ = construir()
    args.salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.salida)
    print(f"Libro creado: {args.salida}")


if __name__ == "__main__":
    main()
